"""XLSX exporter - branded multi-sheet workbook via openpyxl."""

from __future__ import annotations

from pathlib import Path
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.reporting.bundle import ReportBundle

BRAND = {
    "ink": "251B37",
    "accent": "6C5CE7",
    "gold": "FDCB6E",
    "cyan": "74B9FF",
    "green": "00B894",
    "red": "E17055",
    "muted": "6C7A89",
    "bg": "F4F7FB",
}

_SEV_ORDER = ["critical", "high", "medium", "low", "info"]


def _safe(v) -> str:
    if v is None:
        return ""
    return str(v)


def export_xlsx(bundle: ReportBundle, out_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Overview"
    _style_overview(ws, bundle)

    ws2 = wb.create_sheet("Flow Verdicts")
    _style_verdicts(ws2, bundle.verdict_rows())

    ws3 = wb.create_sheet("Findings")
    _style_findings(ws3, bundle.findings)

    ws4 = wb.create_sheet("Framework Mapping")
    _style_frameworks(ws4, bundle.framework_stats)

    ws5 = wb.create_sheet("Raw Records")
    _style_raw(ws5, bundle.summary_of_records())

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))
    return out_path


def _header_row(ws, row: int, headers: List[str], fill: str = "6C5CE7") -> None:
    f = Font(bold=True, color="FFFFFF", size=11)
    fl = PatternFill("solid", fgColor=fill)
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = f
        c.fill = fl
        c.border = Border(bottom=Side(style="thin", color=BRAND["ink"]))
        c.alignment = Alignment(horizontal="center", vertical="center")


def _autofit(ws, widths: List[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _fill_severity(sev: str) -> PatternFill:
    color = {
        "critical": BRAND["red"],
        "high": "E17055",
        "medium": "FDCB6E",
        "low": "74B9FF",
        "info": "B2BEC3",
    }.get(sev, BRAND["muted"])
    return PatternFill("solid", fgColor=color)


def _style_overview(ws, bundle: ReportBundle) -> None:
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:F1")
    ws["A1"] = bundle.project
    ws["A1"].font = Font(bold=True, size=18, color=BRAND["ink"])
    ws.merge_cells("A2:F2")
    ws["A2"] = f"Version {bundle.version}  |  Generated {bundle.generated_at}"
    ws["A2"].font = Font(size=10, color=BRAND["muted"])

    rows = [
        ("Flows analysed", len(bundle.verdicts)),
        ("Findings raised", len(bundle.findings)),
        ("Scenarios", ", ".join(sorted(bundle.meta["scenarios"]))),
        ("Frameworks", "OWASP Top 10 (2021), NIST CSF 2.0, ISO/IEC 27001:2022"),
    ]
    r0 = 4
    for i, (k, v) in enumerate(rows):
        ws.cell(row=r0 + i, column=1, value=k).font = Font(bold=True, color=BRAND["accent"])
        ws.merge_cells(start_row=r0 + i, start_column=2, end_row=r0 + i, end_column=6)
        ws.cell(row=r0 + i, column=2, value=v)
    _autofit(ws, [18, 30, 14, 14, 14, 14])


def _style_verdicts(ws, rows: List[dict]) -> None:
    ws.sheet_view.showGridLines = False
    headers = ["Record", "Scenario", "SNI", "Host", "Dest IP", "CDN", "Score", "Verdict", "Findings"]
    _header_row(ws, 1, headers, BRAND["accent"])
    widths = [10, 14, 26, 26, 16, 14, 8, 12, 8]
    for r_i, r in enumerate(rows, start=2):
        vals = [r["record_id"], r["scenario"], r["sni"], r["host_header"],
                r["dest_ip"], r["cdn_owner"] or "-", r["score"], r["label"], r["findings"]]
        for c_i, v in enumerate(vals, start=1):
            cell = ws.cell(row=r_i, column=c_i, value=v)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        score_cell = ws.cell(row=r_i, column=7)
        score_cell.fill = PatternFill("solid", fgColor=_score_fill(r["score"]))
        score_cell.font = Font(bold=True, color="FFFFFF")
    _autofit(ws, widths)


def _score_fill(score: int) -> str:
    if score >= 70:
        return BRAND["red"]
    if score >= 40:
        return "E17055"
    return BRAND["green"]


def _style_findings(ws, findings) -> None:
    ws.sheet_view.showGridLines = False
    headers = ["Record", "ID", "Severity", "Check", "Title", "Evidence", "Frameworks"]
    _header_row(ws, 1, headers, BRAND["gold"])
    widths = [10, 12, 10, 22, 40, 60, 22]
    for r_i, f in enumerate(findings, start=2):
        vals = [f.record_id.split("-")[0], f.finding_id, f.severity.value, f.check,
                f.title, f.evidence, ", ".join(f.refs)]
        for c_i, v in enumerate(vals, start=1):
            cell = ws.cell(row=r_i, column=c_i, value=v)
            cell.alignment = Alignment(vertical="top", wrap_text=(c_i in (5, 6, 7)))
        sev_cell = ws.cell(row=r_i, column=3)
        sev_cell.fill = _fill_severity(f.severity.value)
        sev_cell.font = Font(bold=True, color="FFFFFF")
        sev_cell.alignment = Alignment(horizontal="center", vertical="center")
    _autofit(ws, widths)


def _style_frameworks(ws, stats: dict) -> None:
    ws.sheet_view.showGridLines = False
    headers = ["Framework", "Controls", "Relevant", "Adopt", "Review", "Monitoring"]
    _header_row(ws, 1, headers, BRAND["cyan"])
    for i, (fw, s) in enumerate(stats.items(), start=2):
        vals = [fw, s["total"], s["relevant"], s["adopt"], s["review"], s["monitoring"]]
        for c_i, v in enumerate(vals, start=1):
            cell = ws.cell(row=i, column=c_i, value=v)
            cell.alignment = Alignment(horizontal="center", vertical="center")
    _autofit(ws, [28, 10, 10, 10, 10, 12])


def _style_raw(ws, rows: List[dict]) -> None:
    ws.sheet_view.showGridLines = False
    headers = list(rows[0].keys()) if rows else ["record_id"]
    _header_row(ws, 1, headers, BRAND["ink"])
    for r_i, r in enumerate(rows, start=2):
        for c_i, h in enumerate(headers, start=1):
            v = r.get(h)
            if isinstance(v, (list, dict)):
                v = ",".join(map(_safe, v))
            ws.cell(row=r_i, column=c_i, value=v)
    _autofit(ws, [12, 14, 16, 10, 16, 10, 26, 26, 24, 12, 32, 24, 20, 10, 16, 10, 20])