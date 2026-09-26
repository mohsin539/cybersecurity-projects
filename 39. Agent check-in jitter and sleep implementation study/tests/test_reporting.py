from __future__ import annotations

import json
import os

from src.reporting import csv_writer
from src.reporting.bundle import write_full_bundle
from src.reporting.report_data import build_report_model
from src.security.audit import AuditLog
from src.security import vault


def _sample_model():
    return {
        "run_id": "abc123qrst",
        "generated_at": "2026-09-20 00:00:00Z",
        "scenario": {"agents": 100, "jitter_pct": 25.0, "sleep_mode": "fixed"},
        "scenario_yaml": "agents: 100",
        "config_hash": "h" * 64,
        "verdict_badge": "PASS",
        "verdict_text": "candidate wins",
        "summary_rows": [
            {"metric": "Herd coefficient", "metric_key": "herd_coef", "candidate": "1.10",
             "candidate_raw": 1.1, "candidate_std": 0.1, "baseline": "12.0", "baseline_raw": 12.0,
             "baseline_std": 1.0},
            {"metric": "Uniformity", "metric_key": "uniformity", "candidate": "0.99",
             "candidate_raw": 0.99, "candidate_std": 0.01},
        ],
        "comparison_rows": [
            {"metric": "Herd coefficient", "candidate_mean": 1.1, "baseline_mean": 12.0,
             "u_stat": 0.0, "u_p": "0.0001", "t_stat": -9.0, "t_p": "0.0001",
             "hedges_g": -2.5, "effect": "large", "direction": "lower_better",
             "verdict": "candidate wins", "significant": "yes"},
        ],
        "replica_rows": [
            {"arm": "candidate", "replica": 1, "herd_coef": 1.1, "uniformity": 0.99},
            {"arm": "baseline", "replica": 1, "herd_coef": 12.0, "uniformity": 0.1},
        ],
        "buckets": {"candidate": [2, 1, 3, 2], "baseline": [20, 0, 0, 20]},
        "engine": "exact",
        "seeds_json": "{}",
    }


def test_csv_injection_guard(tmp_path):
    p = tmp_path / "t.csv"
    csv_writer.write_csv_bundle(_sample_model(), str(tmp_path))
    assert (tmp_path / "summary.csv").exists()
    raw = (tmp_path / "summary.csv").read_text(encoding="utf-8-sig")
    assert "Herd coefficient" in raw


def test_csv_cell_guards_formula():
    from src.security.sanitize import csv_cell
    assert csv_cell("=cmd") == "'=cmd"
    assert csv_cell("-2+3") == "'-2+3"
    assert csv_cell("+SUM(A1)") == "'+SUM(A1)"
    assert csv_cell("normal") == "normal"


def test_html_writer_escapes(tmp_path):
    from src.reporting import html_writer
    model = _sample_model()
    model["verdict_text"] = "<script>alert(1)</script>"
    p = html_writer.write_html(model, str(tmp_path / "report.html"))
    text = open(p, encoding="utf-8").read()
    assert "<script>alert" not in text
    assert "abc123qrst" in text


def test_xlsx_writer_creates_workbook(tmp_path):
    from src.reporting import xlsx_writer
    p = xlsx_writer.write_xlsx(_sample_model(), str(tmp_path / "r.xlsx"))
    assert os.path.exists(p)
    from openpyxl import load_workbook
    wb = load_workbook(p)
    assert "01 Summary" in wb.sheetnames


def test_audit_hash_chain(tmp_path):
    log = AuditLog(str(tmp_path / "audit.log"))
    assert log.verify()
    log.append("test.one", {"a": 1})
    log.append("test.two", {"b": 2})
    assert log.verify()
    with open(str(tmp_path / "audit.log"), "a", encoding="utf-8") as fh:
        fh.write('{"prev":"0000000000000000000000000000000000000000000000000000000000000000","hash":"zzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzzz","payload":{"action":"evil","actor":"x","ts":"2026-01-01","run_id":"0","payload":{}}}\n')
    assert not log.verify()


def test_vault_roundtrip():
    if not vault.crypto_available():
        return
    blob = vault.encrypt_bytes("pw", b"secret-report")
    assert vault.decrypt_bytes("pw", blob) == b"secret-report"
    try:
        vault.decrypt_bytes("wrong", blob)
        assert False, "decrypt with wrong passphrase must raise"
    except Exception:
        pass


def test_bundle_writes_manifest(tmp_path):
    from src.paths import audit_log_path
    audit = AuditLog(str(tmp_path / "a.log"))
    model = _sample_model()
    paths = write_full_bundle(model, str(tmp_path / "out"), audit)
    assert any(k == "sha256.manifest.json" for k in paths)
    manifest = json.load(open(paths["sha256.manifest.json"], encoding="utf-8"))
    assert manifest["scenario.config_hash"] == "h" * 64
    assert audit.verify()