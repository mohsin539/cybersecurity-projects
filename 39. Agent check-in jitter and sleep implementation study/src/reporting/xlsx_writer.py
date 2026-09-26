"""Excel (.xlsx) workbook writer - styled sheets, frozen headers, metrics."""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CYAN_FILL = PatternFill("solid", fgColor="0891B2")
DARK_FILL = PatternFill("solid", fgColor="0F172A")
ALT_FILL = PatternFill("solid", fgColor="112240")
BAND_FILL = PatternFill("solid", fgColor="1E3A5F")
GOOD = PatternFill("solid", fgColor="065F46")
WARN = PatternFill("solid", fgColor="B45309")
BAD = PatternFill("solid", fgColor="991B1B")


def _style_header(ws, row=1, n_cols=None):
    n_cols = n_cols or ws.max_column
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = CYAN_FILL
        cell.alignment = Alignment(vertical="center")


def _autofit(ws, widths=None):
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        return
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        mx = max(len(str(c.value or "")) for c in col) if col else 10
        ws.column_dimensions[letter].width = min(max(mx + 2, 10), 46)


def write_xlsx(model: dict, path: str) -> str:
    wb = Workbook()

    # -- manifest
    ws = wb.active
    ws.title = "00 Manifest"
    ws.append(["Agent Check-in Jitter & Sleep - Study Manifest"])
    ws.append([])
    rows = [
        ("run_id", model["run_id"]),
        ("generated_at", model["generated_at"]),
        ("config_hash", model["config_hash"]),
        ("verdict", f"{model['verdict_badge']} - {model['verdict_text']}"),
        ("engine", model["engine"]),
        ("", ""),
        ("scenario", ""),
    ]
    for k, v in rows:
        ws.append([k, v])
    for k, v in model["scenario"].items():
        ws.append([f"  {k}", v])
    ws.append([])
    ws.append(["seed_registry (replay)", ""])
    for line in model["seeds_json"].splitlines():
        ws.append([line])
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 90
    ws["A1"].font = Font(bold=True, size=14, color="0891B2")

    # -- summary
    ws = wb.create_sheet("01 Summary")
    has_base = bool(model["comparison_rows"]) or (
        bool(model["summary_rows"]) and "baseline" in model["summary_rows"][0]
    )
    if has_base:
        ws.append(["Metric", "Candidate", "Candidate \u00b1 std", "Baseline", "Baseline \u00b1 std"])
    else:
        ws.append(["Metric", "Candidate", "Candidate \u00b1 std"])
    for r in model["summary_rows"]:
        if "baseline" in r:
            ws.append([r["metric"], r["candidate"], f"{r['candidate_raw']} \u00b1 {r['candidate_std']}",
                       r["baseline"], f"{r['baseline_raw']} \u00b1 {r['baseline_std']}"])
        else:
            ws.append([r["metric"], r["candidate"], f"{r['candidate_raw']} \u00b1 {r['candidate_std']}"])
    _style_header(ws)
    ws.freeze_panes = "A2"
    _autofit(ws)

    # -- comparison
    ws = wb.create_sheet("02 Comparisons")
    if model["comparison_rows"]:
        ws.append(["Metric", "Candidate mean", "Baseline mean", "M-W U", "M-W p(adj)", "Welch t", "t p", "Hedges g", "Effect", "Verdict", "Significant"])
        for r in model["comparison_rows"]:
            ws.append([r["metric"], r["candidate_mean"], r["baseline_mean"], r["u_stat"], r["u_p"], r["t_stat"], r["t_p"], r["hedges_g"], r["effect"], r["verdict"], r["significant"]])
        _style_header(ws)
        ws.freeze_panes = "A2"
        _autofit(ws)
    else:
        ws.append(["No comparison arm configured."])
        _autofit(ws)

    # -- per replica
    ws = wb.create_sheet("03 Per-Replica")
    if model["replica_rows"]:
        keys = list(model["replica_rows"][0].keys())
        ws.append(keys)
        for r in model["replica_rows"]:
            ws.append([r.get(k, "") for k in keys])
        _style_header(ws)
        ws.freeze_panes = "A2"
        _autofit(ws)

    # -- buckets (first replica of each arm)
    ws = wb.create_sheet("04 Buckets")
    buck = model["buckets"]
    n = max(len(buck.get("candidate", [])), len(buck.get("baseline", [])))
    ws.append(["second"] + (["candidate"] if buck.get("baseline") else ["candidate", "baseline"]))
    pad = lambda arr: arr + [""] * (n - len(arr))  # noqa: E731
    for i in range(n):
        row = [i]
        row.append(pad(buck["candidate"])[i] if buck.get("candidate") else "")
        if buck.get("baseline"):
            row.append(pad(buck["baseline"])[i])
        ws.append(row)
    _style_header(ws)
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16

    wb.save(path)
    return path