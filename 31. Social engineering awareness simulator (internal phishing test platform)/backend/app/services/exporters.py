import csv
import html as html_lib
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RISK_FILLS = {
    "critical": PatternFill("solid", fgColor="C0504D"),
    "high": PatternFill("solid", fgColor="F9CD9C"),
    "medium": PatternFill("solid", fgColor="FFE699"),
    "low": PatternFill("solid", fgColor="C6E0B4"),
}
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _flagged(score: float) -> str:
    if score >= 70:
        return "critical"
    if score >= 40:
        return "high"
    if score >= 15:
        return "medium"
    return "low"


def _style_sheet(ws, headers, widths):
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(col)].width = widths.get(col, 18)
    ws.freeze_panes = "A2"


def build_xlsx(data: dict, path: str) -> int:
    wb = Workbook()

    ws = wb.active
    ws.title = "Executive Summary"
    ws.append(["Metric", "Value"])
    ws.append(["Generated", data["generated_at"]])
    ws.append(["Total campaigns", data["summary"]["campaigns"]])
    ws.append(["Total delivered", data["summary"]["sent"]])
    ws.append(["Opened", data["summary"]["opened"]])
    ws.append(["Clicked", data["summary"]["clicked"]])
    ws.append(["Submitted credentials", data["summary"]["submitted"]])
    ws.append(["Reported to SOC", data["summary"]["reported"]])
    ws.append(["Average SE-Index", data["summary"]["avg_se_index"]])
    _style_sheet(ws, ["Metric", "Value"], {1: 30, 2: 24})

    ws2 = wb.create_sheet("Campaign Performance")
    headers = ["Campaign", "Vector", "Status", "Sent", "Opened", "Clicked", "Submitted", "Reported"]
    ws2.append(headers)
    for c in data["campaigns"]:
        ws2.append(
            [
                c["name"],
                c["vector"],
                c["status"],
                c["sent"],
                c["opened"],
                c["clicked"],
                c["submitted"],
                c["reported"],
            ]
        )
    _style_sheet(ws2, headers, {1: 24, 3: 12, 4: 9, 5: 9, 6: 9, 7: 9, 8: 9})

    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Campaign Behaviour"
    data_ref = Reference(ws2, min_col=4, min_row=1, max_col=8, max_row=min(ws2.max_row, 15))
    cats_ref = Reference(ws2, min_col=1, min_row=2, max_row=min(ws2.max_row, 15))
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    chart.dataLabels = DataLabelList()
    ws2.add_chart(chart, "J2")

    ws3 = wb.create_sheet("Employee Risk Matrix")
    headers3 = ["Employee", "Branch", "SE-Index", "Clicks", "Submissions", "Training OK", "Flag"]
    ws3.append(headers3)
    for r in data["risk_rows"]:
        flag = _flagged(r["se_index"])
        ws3.append(
            [
                r["full_name"],
                r["branch"],
                r["se_index"],
                r["clicks"],
                r["submissions"],
                "Yes" if r["trainings_ok"] else "No",
                flag.upper(),
            ]
        )
        risk_row = ws3.max_row
        ws3.cell(row=risk_row, column=3).fill = RISK_FILLS.get(flag)
        ws3.cell(row=risk_row, column=7).fill = RISK_FILLS.get(flag)
    _style_sheet(ws3, headers3, {1: 22, 4: 9, 5: 12, 6: 11, 7: 10})

    ws4 = wb.create_sheet("Raw Events")
    headers4 = ["Time", "Employee", "Campaign", "Event"]
    ws4.append(headers4)
    for e in data["events"][:5000]:
        ws4.append(e)
    _style_sheet(
        ws4,
        headers4,
        {1: 22, 2: 22, 3: 22, 4: 16},
    )

    wb.save(path)
    return os.path.getsize(path)


def build_csv(data: dict, path: str) -> int:
    rows = [["Social Engineering Awareness Simulator - Report"]]
    rows.append(["Generated", data["generated_at"]])
    rows.append(["Campaigns", data["summary"]["campaigns"]])
    rows.append(["Delivered", data["summary"]["sent"]])
    rows.append(["Opened", data["summary"]["opened"]])
    rows.append(["Clicked", data["summary"]["clicked"]])
    rows.append(["Submitted", data["summary"]["submitted"]])
    rows.append(["Reported", data["summary"]["reported"]])
    rows.append(["Avg SE-Index", data["summary"]["avg_se_index"]])
    rows.append([])
    rows.append(["Employee", "Branch", "SE-Index", "Clicks", "Submissions", "Training OK"])
    for r in data["risk_rows"]:
        rows.append(
            [
                r["full_name"],
                r["branch"],
                r["se_index"],
                r["clicks"],
                r["submissions"],
                "Yes" if r["trainings_ok"] else "No",
            ]
        )

    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh, quoting=csv.QUOTE_MINIMAL)
        writer.writerows(rows)
    return os.path.getsize(path)


def build_html(data: dict, path: str) -> int:
    esc = html_lib.escape

    def flag_pill(score: float) -> str:
        fl = _flagged(score)
        colors = {"critical": "C0504D", "high": "E36C09", "medium": "BF8F00", "low": "548235"}
        return (
            f'<span style="background:#{colors[fl]};color:#fff;padding:2px 10px;'
            f'border-radius:999px;font-size:12px;font-weight:700;">{fl.upper()}</span>'
        )

    campaign_rows = "".join(
        f"<tr><td>{esc(c['name'])}</td><td>{esc(c['vector'])}</td><td>{esc(c['status'])}</td>"
        f"<td>{c['sent']}</td><td>{c['opened']}</td><td>{c['clicked']}</td>"
        f"<td>{c['submitted']}</td><td>{c['reported']}</td></tr>"
        for c in data["campaigns"]
    )

    risk_rows = "".join(
        f"<tr><td>{esc(r['full_name'])}</td><td>{esc(r['branch'])}</td>"
        f"<td>{r['se_index']}</td>{flag_pill(r['se_index'])}"
        f"<td>{r['clicks']}</td><td>{r['submissions']}</td>"
        f"<td>{'Yes' if r['trainings_ok'] else 'No'}</td></tr>"
        for r in data["risk_rows"]
    )

    event_rows = "".join(
        f"<tr><td>{esc(e[0])}</td><td>{esc(e[1])}</td><td>{esc(e[2])}</td><td>{esc(e[3])}</td></tr>"
        for e in data["events"][:300]
    )

    m = data["summary"]
    kpis = "".join(
        f'<div style="flex:1;background:linear-gradient(135deg,#0e7490,#134e4a);color:#fff;'
        f'border-radius:14px;padding:16px;text-align:center;margin:6px;">'
        f'<div style="font-size:26px;font-weight:800;">{v}</div>'
        f'<div style="font-size:12px;opacity:.85;">{k}</div></div>'
        for k, v in [
            ("Campaigns", m["campaigns"]),
            ("Delivered", m["sent"]),
            ("Clicked", m["clicked"]),
            ("Avg SE-Index", m["avg_se_index"]),
        ]
    )

    page = f"""<!DOCTYPE html>
<html lang="en">
<head><meta charset="utf-8"><title>SEAS Dashboard Report</title>
<style>
  body{{font-family:'Segoe UI',Roboto,Arial,sans-serif;color:#1a2433;background:#f4f7fb;margin:0;padding:24px;}}
  .hero{{background:linear-gradient(120deg,#0a1a33,#0d3b52);color:#fff;border-radius:16px;padding:26px;}}
  .hero h1{{margin:0 0 6px;}}
  .card{{background:#fff;border:1px solid #dbe4f0;border-radius:14px;padding:18px;margin:16px 0;}}
  table{{border-collapse:collapse;width:100%;font-size:13px;}}
  th,td{{border:1px solid #dbe4f0;padding:8px 10px;text-align:left;}}
  th{{background:#1f4e79;color:#fff;}}
  .kpis{{display:flex;flex-wrap:wrap;gap:10px;margin-top:14px;}}
  .muted{{color:#5b6b80;}}
  @media print{{body{{background:#fff;}}}}
</style></head>
<body>
  <div class="hero">
    <h1>Social Engineering Awareness Simulator</h1>
    <div class="muted" style="color:#bfd9ff;">Generated {esc(data['generated_at'])} &middot; Internal phishing test platform report</div>
  </div>
  <div class="kpis">{kpis}</div>
  <div class="card"><h2>Campaign Performance</h2>
    <table><thead><tr><th>Campaign</th><th>Vector</th><th>Status</th><th>Sent</th><th>Opened</th><th>Clicked</th><th>Submitted</th><th>Reported</th></tr></thead>
    <tbody>{campaign_rows}</tbody></table></div>
  <div class="card"><h2>Employee Risk Matrix (SE-Index)</h2>
    <table><thead><tr><th>Employee</th><th>Branch</th><th>SE-Index</th><th>Flag</th><th>Clicks</th><th>Submissions</th><th>Training OK</th></tr></thead>
    <tbody>{risk_rows}</tbody></table></div>
  <div class="card"><h2>Raw Event Chronicle</h2>
    <table><thead><tr><th>Time</th><th>Employee</th><th>Campaign</th><th>Event</th></tr></thead>
    <tbody>{event_rows}</tbody></table></div>
  <script>window.print=function(){{}};</script>
</body></html>
"""

    with open(path, "w", encoding="utf-8") as fh:
        fh.write(page)
    return os.path.getsize(path)


def build_export(data: dict, fmt: str, path: str) -> int:
    builders = {"xlsx": build_xlsx, "csv": build_csv, "html": build_html}
    return builders[fmt](data, path)