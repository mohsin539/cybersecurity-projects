"""Reporting Engine (architecture.md §9) — PDF, DOCX, XLSX, CSV, JSON + evidence ZIP."""
import csv
import io
import json
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from fastapi.responses import Response
from sqlalchemy.orm import Session

from ..models import (
    Assessment, AssessmentDecision, Control, ControlMapping, Evidence, EvidenceLink,
    Framework, RemediationTicket, RiskPoint, ScanResult,
)

EXPORT_DIR = Path(__file__).resolve().parent.parent.parent / "data" / "exports"
EXPORT_DIR.mkdir(parents=True, exist_ok=True)

CONTENT_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv",
    "json": "application/json",
    "zip": "application/zip",
}


def _iso(dt):
    return dt.isoformat() if hasattr(dt, "isoformat") else str(dt)


def _control_matrix_rows(db: Session, framework_code: str = None) -> list[dict]:
    """Rows = control details + latest decision status + evidence count."""
    rows = []
    fws = db.query(Framework).filter(Framework.is_active.is_(True)).all()
    if framework_code:
        fws = [f for f in fws if f.code == framework_code]
    for fw in fws:
        latest = (
            db.query(Assessment)
            .filter(Assessment.framework_code == fw.code, Assessment.status.in_(["COMPLETED", "REVIEWED"]))
            .order_by(Assessment.completed_at.desc()).first()
        )
        status_map = {}
        if latest:
            for d in db.query(AssessmentDecision).filter(AssessmentDecision.assessment_id == latest.id).all():
                status_map[d.control_id] = d.status
        for ctrl in db.query(Control).filter(Control.framework_id == fw.id).order_by(Control.code).all():
            st = status_map.get(ctrl.id, ctrl.implementation_status if ctrl.implementation_status in
                                ("COMPLIANT", "PARTIAL", "NON_COMPLIANT", "NOT_ASSESSED", "NOT_APPLICABLE")
                                else "NOT_ASSESSED")
            ev_count = db.query(EvidenceLink).filter(EvidenceLink.control_id == ctrl.id).count()
            mapped = db.query(ControlMapping).filter(ControlMapping.source_id == ctrl.id).count()
            rows.append({
                "framework": fw.code, "control_code": ctrl.code, "title": ctrl.title,
                "category": ctrl.category, "intent": ctrl.intent, "status": st,
                "evidence_count": ev_count, "mapped_to": mapped, "owner": ctrl.owner or "",
            })
    return rows


def _risk_rows(db: Session) -> list[dict]:
    return [{
        "risk_id": r.id, "title": r.title, "tier": r.tier or "N/A",
        "likelihood": r.likelihood, "impact": r.impact, "cvss": r.cvss,
        "raw_score": r.raw_score, "residual_score": r.residual_score, "status": r.status,
    } for r in db.query(RiskPoint).all()]


def _assessment_rows(db: Session) -> list[dict]:
    return [{
        "id": a.id, "name": a.name, "framework": a.framework_code, "method": a.method,
        "status": a.status, "score": a.result_score, "findings": a.findings_count,
        "completed_at": _iso(a.completed_at), "created_by": a.created_by,
    } for a in db.query(Assessment).order_by(Assessment.id.desc()).all()]


def _evidence_rows(db: Session) -> list[dict]:
    rows = []
    for ev in db.query(Evidence).order_by(Evidence.id.desc()).all():
        links = db.query(EvidenceLink).filter(EvidenceLink.evidence_id == ev.id).count()
        rows.append({
            "id": ev.id, "title": ev.title, "type": ev.artefact_type,
            "source": ev.source_system, "sha256": ev.sha256, "chain_hash": ev.chain_hash,
            "control_links": links, "uploaded_by": ev.uploaded_by,
            "created_at": _iso(ev.created_at), "worm_locked": ev.worm_locked,
        })
    return rows


def _scan_rows(db: Session) -> list[dict]:
    rows = []
    for s in db.query(ScanResult).order_by(ScanResult.scanned_at.desc()).limit(500).all():
        rows.append({
            "asset": s.asset.name if s.asset else "", "scanner": s.scanner,
            "finding_type": s.finding_type, "title": s.title, "severity": s.severity,
            "cvss": s.cvss, "status": s.status, "scanned_at": _iso(s.scanned_at),
        })
    return rows


def _mapping_rows(db: Session) -> list[dict]:
    rows = []
    for m in db.query(ControlMapping).all():
        src, tgt = m.source, m.target
        src_fw = db.query(Framework).filter(Framework.id == src.framework_id).first()
        tgt_fw = db.query(Framework).filter(Framework.id == tgt.framework_id).first()
        rows.append({
            "from_framework": src_fw.code if src_fw else "", "from_control": src.code,
            "to_framework": tgt_fw.code if tgt_fw else "", "to_control": tgt.code,
            "map_type": m.map_type, "rationale": m.rationale or "",
        })
    return rows


# ---------------------------------------------------------------------------
# Bangladesh Bank ICT 2015 — Quarterly F&R Return (architecture.md §9)
# ---------------------------------------------------------------------------
def _latest_bb_assessment(db: Session):
    return (
        db.query(Assessment)
        .filter(Assessment.framework_code == "BBICT2015",
                Assessment.status.in_(["COMPLETED", "REVIEWED"]))
        .order_by(Assessment.completed_at.desc()).first()
    )


def _bb_status_map(db: Session) -> tuple[dict, Assessment | None]:
    latest = _latest_bb_assessment(db)
    status_map = {}
    if latest:
        for d in db.query(AssessmentDecision).filter(AssessmentDecision.assessment_id == latest.id).all():
            status_map[d.control_id] = d.status
    return status_map, latest


def _overlap_codes(db: Session, control_id: int) -> tuple[str, str, str]:
    """ISO / NIST / OWASP overlap codes for a BB control (reverse of seeded edges)."""
    iso, nist, owasp = [], [], []
    for m in db.query(ControlMapping).filter(ControlMapping.target_id == control_id).all():
        ctrl = m.source
        fw = db.query(Framework).filter(Framework.id == ctrl.framework_id).first()
        if fw and fw.code == "ISO27001":
            iso.append(ctrl.code)
        elif fw and fw.code == "NISTCSF":
            nist.append(ctrl.code)
        elif fw and fw.code == "OWASP2021":
            owasp.append(ctrl.code)
    return ",".join(iso), ",".join(nist), ",".join(owasp)


def _fr_chapter_rows(db: Session) -> list[dict]:
    fw = db.query(Framework).filter(Framework.code == "BBICT2015").first()
    if not fw:
        return []
    status_map, _ = _bb_status_map(db)
    chapters = {}
    for ctrl in db.query(Control).filter(Control.framework_id == fw.id).order_by(Control.code).all():
        chapters.setdefault(ctrl.code, {"title": ctrl.title, "statuses": []})
        chapters[ctrl.code]["statuses"].append(status_map.get(ctrl.id, "NOT_ASSESSED"))
    rows = []
    for code, ch in chapters.items():
        counts = Counter(ch["statuses"])
        total = len(ch["statuses"]) or 1
        weighted = (counts.get("COMPLIANT", 0) + 0.5 * counts.get("PARTIAL", 0)) / total * 100
        rows.append({
            "bb_chapter": code, "requirement": ch["title"],
            "controls_total": len(ch["statuses"]),
            "compliant": counts.get("COMPLIANT", 0),
            "partial": counts.get("PARTIAL", 0),
            "non_compliant": counts.get("NON_COMPLIANT", 0),
            "not_assessed": counts.get("NOT_ASSESSED", 0),
            "maturity_pct": round(weighted, 1),
        })
    return rows


def _fr_control_rows(db: Session) -> list[dict]:
    fw = db.query(Framework).filter(Framework.code == "BBICT2015").first()
    if not fw:
        return []
    status_map, _ = _bb_status_map(db)
    rows = []
    seq = 1
    for ctrl in db.query(Control).filter(Control.framework_id == fw.id).order_by(Control.code).all():
        links = db.query(EvidenceLink).filter(EvidenceLink.control_id == ctrl.id).all()
        iso, nist, owasp = _overlap_codes(db, ctrl.id)
        rows.append({
            "seq": seq, "bb_chapter": ctrl.code, "requirement": ctrl.title,
            "category": ctrl.category or "", "status": status_map.get(ctrl.id, "NOT_ASSESSED"),
            "evidence_count": len(links),
            "evidence_refs": "; ".join(f"ev-{l.evidence_id}" for l in links) or "—",
            "iso_overlap": iso, "nist_overlap": nist, "owasp_overlap": owasp,
        })
        seq += 1
    return rows


def _fr_evidence_rows(db: Session) -> list[dict]:
    rows = []
    for ev in db.query(Evidence).order_by(Evidence.id).all():
        linked = []
        for link in db.query(EvidenceLink).filter(EvidenceLink.evidence_id == ev.id).all():
            ctrl = link.control
            if ctrl:
                fw = db.query(Framework).filter(Framework.id == ctrl.framework_id).first()
                linked.append(f"{fw.code}:{ctrl.code}" if fw else ctrl.code)
        rows.append({
            "evidence_id": ev.id, "title": ev.title, "artefact_type": ev.artefact_type,
            "source_system": ev.source_system, "sha256": ev.sha256,
            "worm_locked": ev.worm_locked, "control_links": "; ".join(linked),
        })
    return rows


def _fr_findings_rows(db: Session) -> list[dict]:
    rows = []
    for r in db.query(RiskPoint).filter(RiskPoint.status != "RESOLVED").all():
        ticket = db.query(RemediationTicket).filter(RemediationTicket.risk_id == r.id).first()
        rows.append({
            "risk_id": r.id, "title": r.title, "tier": r.tier or "",
            "cvss": r.cvss, "status": r.status,
            "ticket": f"#T{ticket.id}" if ticket else "—",
            "ticket_priority": ticket.priority if ticket else "—",
            "due_at": _iso(ticket.due_at) if ticket and ticket.due_at else "—",
        })
    return rows


def _fr_attestation_rows(db: Session, actor: str = None, period: str = None) -> list[dict]:
    now = datetime.now(timezone.utc)
    return [
        {"field": "Return Type", "value": "Bangladesh Bank ICT Guidelines 2015 — Quarterly Financial & Regulatory (F&R) Return"},
        {"field": "Reporting Period", "value": period or f"Q{(now.month - 1) // 3 + 1} {now.year}"},
        {"field": "Reporting Institution", "value": "Compliance Automation Suite — Demo Bank (FinTech)"},
        {"field": "Return Generated By", "value": actor or "system"},
        {"field": "Return Generated At", "value": now.isoformat()},
        {"field": "ISMS Status", "value": "Active — ISO/IEC 27001:2022 ISMS in continuous compliance"},
        {"field": "Statement",
         "value": "Certified that the controls declared herein reflect the verified evidence-and-assessment "
                  "state of the institution at the time of generation. Evidence artefacts are WORM-locked and "
                  "SHA-256 hash-chained (ISO A.8.15 / BB Ch-14)."},
        {"field": "CISO Attestation", "value": "Name: ______________________    Signature: ____________    Date: __________"},
    ]


def _bb_fr_sheets(db: Session, actor: str = None, period: str = None) -> dict[str, list[dict]]:
    return {
        "Cover Letter": _fr_attestation_rows(db, actor, period),
        "Chapter Summary": _fr_chapter_rows(db),
        "Control Return": _fr_control_rows(db),
        "Evidence Index": _fr_evidence_rows(db),
        "Findings & Remediation": _fr_findings_rows(db),
    }


def _bb_fr_sections(db: Session, actor: str = None, period: str = None) -> list[tuple[str, list[dict]]]:
    return [
        ("Cover Letter", _fr_attestation_rows(db, actor, period)),
        ("Chapter Summary", _fr_chapter_rows(db)),
        ("Control Return", _fr_control_rows(db)),
        ("Evidence Index", _fr_evidence_rows(db)),
        ("Findings & Remediation", _fr_findings_rows(db)),
    ]


def build_bb_fr_return(db: Session, fmt: str, actor: str = None,
                       period: str = None) -> Response:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename_base = f"bb_fr_return_{stamp}"
    if fmt == "csv":
        return Response(_csv(_fr_control_rows(db)), media_type=CONTENT_TYPES["csv"],
                        headers={"Content-Disposition": f"attachment; filename={filename_base}.csv"})
    if fmt == "xlsx":
        return Response(_xlsx(_bb_fr_sheets(db, actor=actor, period=period)),
                        media_type=CONTENT_TYPES["xlsx"],
                        headers={"Content-Disposition": f"attachment; filename={filename_base}.xlsx"})
    if fmt == "json":
        payload = {name: rows for name, rows in _bb_fr_sections(db, actor=actor, period=period)}
        return Response(_json(payload), media_type=CONTENT_TYPES["json"],
                        headers={"Content-Disposition": f"attachment; filename={filename_base}.json"})
    if fmt == "docx":
        payload = _docx("BB ICT GUIDELINES 2015 — QUARTERLY F&R RETURN",
                        "Compliance Automation Suite — Bangladesh Bank Financial & Regulatory Return",
                        _bb_fr_sections(db, actor=actor, period=period))
        return Response(payload, media_type=CONTENT_TYPES["docx"],
                        headers={"Content-Disposition": f"attachment; filename={filename_base}.docx"})
    if fmt == "pdf":
        payload = _pdf("BB ICT GUIDELINES 2015 — QUARTERLY F&R RETURN",
                       "Compliance Automation Suite — Bangladesh Bank Financial & Regulatory Return",
                       _bb_fr_sections(db, actor=actor, period=period))
        return Response(payload, media_type=CONTENT_TYPES["pdf"],
                        headers={"Content-Disposition": f"attachment; filename={filename_base}.pdf"})
    raise ValueError(f"Unsupported format {fmt}")


# ---------------------------------------------------------------------------
# Generators
# ---------------------------------------------------------------------------
def _csv(data: list[dict]) -> str:
    buf = io.StringIO()
    if not data:
        return "no_data\n"
    writer = csv.DictWriter(buf, fieldnames=list(data[0].keys()))
    writer.writeheader()
    writer.writerows(data)
    return buf.getvalue()


def _json(data, pretty=True) -> str:
    return json.dumps(data, indent=2 if pretty else None, default=str)


def _xlsx(sheets: dict[str, list[dict]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for sheet_name, rows in sheets.items():
        ws = wb.create_sheet(sheet_name[:31])
        if rows:
            headers = list(rows[0].keys())
            ws.append(headers)
            for col, h in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")
            for r in rows:
                ws.append([r.get(h, "") for h in headers])
            for col in range(1, len(headers) + 1):
                width = max(len(str(headers[col - 1])) * 1.5 + 2,
                            *(len(str(r.get(headers[col - 1], ""))) * 1.1 + 2 for r in rows[:200]))
                ws.column_dimensions[get_column_letter(col)].width = min(width, 50)
            ws.auto_filter.ref = ws.dimensions
            ws.freeze_panes = "A2"
        else:
            ws.append(["no_data"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx(title: str, subtitle: str, sections: list[tuple[str, list[dict]]]) -> bytes:
    from docx import Document
    from docx.shared import RGBColor, Pt
    from docx.enum.table import WD_TABLE_ALIGNMENT

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    heading = doc.add_heading(title, level=0)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)

    p = doc.add_paragraph(subtitle)
    p.runs[0].italic = True

    doc.add_paragraph(f"Generated at: {datetime.now(timezone.utc).isoformat()}")

    for sec_title, rows in sections:
        doc.add_heading(sec_title, level=1)
        if not rows:
            doc.add_paragraph("(no data)")
            continue
        headers = list(rows[0].keys())
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Light Grid Accent 1"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = table.rows[0].cells
        for i, h in enumerate(headers):
            hdr[i].text = h
        for r in rows[:500]:
            cells = table.add_row().cells
            for i, h in enumerate(headers):
                cells[i].text = str(r.get(h, ""))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pdf(title: str, subtitle: str, sections: list[tuple[str, list[dict]]]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak,
    )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), rightMargin=28, leftMargin=28,
                            topMargin=28, bottomMargin=28)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitleT", parent=styles["Title"], textColor=colors.HexColor("#1F4E78"))
    sub_style = ParagraphStyle("Sub", parent=styles["Italic"], textColor=colors.HexColor("#5A6772"))
    h1_style = ParagraphStyle("H1c", parent=styles["Heading2"], textColor=colors.HexColor("#0D47A1"))
    body = styles["BodyText"]

    flow = [Paragraph(title, title_style), Paragraph(subtitle, sub_style),
            Spacer(1, 8), Paragraph(f"Generated {datetime.now(timezone.utc).isoformat()}", body),
            Spacer(1, 4)]

    for sec_title, rows in sections:
        flow.append(Paragraph(sec_title, h1_style))
        flow.append(Spacer(1, 4))
        if not rows:
            flow.append(Paragraph("(no data)", body))
            flow.append(Spacer(1, 10))
            continue
        headers = list(rows[0].keys())
        table_data = [[h.upper() for h in headers]]
        for r in rows[:400]:
            table_data.append([str(r.get(h, ""))[:40] for h in headers])
        tbl = Table(table_data, repeatRows=1)
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#EDF2F9"), colors.white]),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B7C6D9")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        flow.append(tbl)
        flow.append(Spacer(1, 12))
    doc.build(flow)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Public entry
# ---------------------------------------------------------------------------
def build_report(db: Session, report_type: str, fmt: str, framework_code: str = None,
                 actor: str = None, period: str = None) -> Response:
    if report_type == "bb_fr":
        return build_bb_fr_return(db, fmt, actor=actor, period=period)
    sections = []
    if report_type == "soc":
        sections = [
            ("Control Matrix", _control_matrix_rows(db, framework_code)),
            ("Risk Register", _risk_rows(db)),
            ("Assessments", _assessment_rows(db)),
            ("Evidence Vault", _evidence_rows(db)),
            ("Open Findings", _scan_rows(db)),
            ("Framework Mappings", _mapping_rows(db)),
        ]
    elif report_type == "controls":
        sections = [("Control Matrix", _control_matrix_rows(db, framework_code))]
    elif report_type == "risk":
        sections = [("Risk Register", _risk_rows(db))]
    elif report_type == "evidence":
        sections = [("Evidence Vault", _evidence_rows(db))]
    elif report_type == "scans":
        sections = [("Open Findings", _scan_rows(db))]
    elif report_type == "mapping":
        sections = [("Framework Mappings", _mapping_rows(db))]
    else:
        raise ValueError(f"Unknown report type {report_type}")

    filename = f"{report_type}_{framework_code or 'all'}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.{fmt}"
    if fmt == "csv":
        raise NotImplementedError("csv uses single-sheet; use /reports/{type}/csv")
    if fmt == "json":
        payload = {name: rows for name, rows in sections}
        return Response(_json(payload), media_type=CONTENT_TYPES["json"],
                        headers={"Content-Disposition": f"attachment; filename={filename}"})
    if fmt == "xlsx":
        payload = {name[:31]: rows for name, rows in sections}
        return Response(_xlsx(payload), media_type=CONTENT_TYPES["xlsx"],
                        headers={"Content-Disposition": f"attachment; filename={filename}"})
    if fmt == "docx":
        payload = _docx(filename.removesuffix(".docx").upper(), "Compliance Automation Suite", sections or [("Report", [])])
        return Response(payload, media_type=CONTENT_TYPES["docx"],
                        headers={"Content-Disposition": f"attachment; filename={filename}"})
    if fmt == "pdf":
        payload = _pdf(filename.removesuffix(".pdf").upper(), "Compliance Automation Suite", sections)
        return Response(payload, media_type=CONTENT_TYPES["pdf"],
                        headers={"Content-Disposition": f"attachment; filename={filename}"})
    raise ValueError(f"Unsupported format {fmt}")


def build_csv(db: Session, report_type: str, framework_code: str = None) -> Response:
    if report_type == "controls":
        data = _control_matrix_rows(db, framework_code)
    elif report_type == "risk":
        data = _risk_rows(db)
    elif report_type == "evidence":
        data = _evidence_rows(db)
    elif report_type == "scans":
        data = _scan_rows(db)
    elif report_type == "mapping":
        data = _mapping_rows(db)
    elif report_type == "assessments":
        data = _assessment_rows(db)
    elif report_type == "bb_fr":
        data = _fr_control_rows(db)
    else:
        raise ValueError(f"Unknown report type {report_type}")
    filename = f"{report_type}_{framework_code or 'all'}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    return Response(_csv(data), media_type=CONTENT_TYPES["csv"],
                    headers={"Content-Disposition": f"attachment; filename={filename}"})


def build_evidence_pack_zip(db: Session, evidence_ids: list[int]) -> Response:
    import zipfile as zf_lib
    buf = io.BytesIO()
    with zf_lib.ZipFile(buf, "w", zf_lib.ZIP_DEFLATED) as zf:
        manifest = {"generated_at": datetime.now(timezone.utc).isoformat(), "artefacts": []}
        for ev in db.query(Evidence).filter(Evidence.id.in_(evidence_ids)).all():
            manifest["artefacts"].append({
                "id": ev.id, "title": ev.title, "sha256": ev.sha256,
                "chain_hash": ev.chain_hash, "mime": ev.mime_type,
                "created_at": _iso(ev.created_at),
            })
            zf.writestr(f"evidence/ev-{ev.id}.desc.txt",
                        f"{ev.title}\n{ev.description or ''}\nsha256: {ev.sha256}")
        zf.writestr("MANIFEST.json", json.dumps(manifest, indent=2, default=str))
    filename = f"evidence_pack_{datetime.now().strftime('%Y%m%d_%H%M%S')}.zip"
    return Response(buf.getvalue(), media_type=CONTENT_TYPES["zip"],
                    headers={"Content-Disposition": f"attachment; filename={filename}"})